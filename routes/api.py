from datetime import datetime, date, timedelta
from flask import Blueprint, jsonify, request, session, abort
from openpyxl import load_workbook
import re

from services.db_service import db
from services.models import Lesson, User, UserRole, Participant, Subscriber
from services.telegram_sender import notify_lesson_change, send_message

api_bp = Blueprint("api", __name__)


def require_admin():
    user_id = session.get("user_id")
    user = User.query.get(user_id) if user_id else None
    if not user or user.role != UserRole.ADMIN or (user.username != "adrkaaa" and user.email != "adrkaaa"):
        abort(403)
    return user


@api_bp.get("/lessons")
def get_lessons():
    month = request.args.get("month", type=int)
    year = request.args.get("year", type=int)
    start_str = request.args.get("start")
    end_str = request.args.get("end")

    query = Lesson.query
    if start_str and end_str:
        try:
            start = datetime.strptime(start_str, "%Y-%m-%d").date()
            end = datetime.strptime(end_str, "%Y-%m-%d").date()
        except Exception:
            return jsonify({"ok": False, "error": "Неверный формат дат"}), 400
        query = query.filter(Lesson.lesson_date >= start, Lesson.lesson_date <= end)
    elif month and year:
        start = date(year, month, 1)
        if month == 12:
            end = date(year + 1, 1, 1)
        else:
            end = date(year, month + 1, 1)
        query = query.filter(Lesson.lesson_date >= start, Lesson.lesson_date < end)

    lessons = query.order_by(Lesson.lesson_date, Lesson.pair).all()
    return jsonify([l.to_dict() for l in lessons])


@api_bp.post("/lessons")
def create_lesson():
    require_admin()
    data = request.get_json(silent=True) or {}

    try:
        lesson_date = datetime.strptime(data.get("date"), "%Y-%m-%d").date()
    except Exception:
        return jsonify({"ok": False, "error": "Некорректная дата"}), 400

    try:
        pair = int(data.get("pair"))
    except Exception:
        return jsonify({"ok": False, "error": "Некорректный номер пары"}), 400

    subject = (data.get("subject") or "").strip()
    teacher = (data.get("teacher") or "").strip()
    audience = (data.get("audience") or "").strip() or None
    status = (data.get("status") or "scheduled").strip()
    comment = (data.get("comment") or "").strip() or None

    if not subject or not teacher:
        return jsonify({"ok": False, "error": "Заполните дисциплину и преподавателя"}), 400

    lesson = Lesson(
        lesson_date=lesson_date,
        pair=pair,
        subject=subject,
        teacher=teacher,
        audience=audience,
        status=status,
        comment=comment,
    )
    db.session.add(lesson)
    db.session.commit()
    notify_lesson_change("created", lesson, old=None)
    return jsonify({"ok": True, "lesson": lesson.to_dict(), "id": lesson.id})


@api_bp.patch("/lessons/<int:lesson_id>")
def update_lesson(lesson_id):
    require_admin()
    lesson = Lesson.query.get_or_404(lesson_id)
    data = request.get_json(silent=True) or {}
    old_data = lesson.to_dict()

    if "date" in data:
        try:
            lesson.lesson_date = datetime.strptime(data.get("date"), "%Y-%m-%d").date()
        except Exception:
            return jsonify({"ok": False, "error": "Некорректная дата"}), 400

    if "pair" in data:
        try:
            lesson.pair = int(data.get("pair"))
        except Exception:
            return jsonify({"ok": False, "error": "Некорректный номер пары"}), 400

    if "subject" in data:
        subject = (data.get("subject") or "").strip()
        if not subject:
            return jsonify({"ok": False, "error": "Дисциплина обязательна"}), 400
        lesson.subject = subject

    if "teacher" in data:
        teacher = (data.get("teacher") or "").strip()
        if not teacher:
            return jsonify({"ok": False, "error": "Преподаватель обязателен"}), 400
        lesson.teacher = teacher

    if "audience" in data:
        lesson.audience = (data.get("audience") or "").strip() or None

    if "status" in data:
        lesson.status = (data.get("status") or "scheduled").strip()

    if "comment" in data:
        lesson.comment = (data.get("comment") or "").strip() or None

    db.session.commit()
    notify_lesson_change("updated", lesson, old=old_data)
    return jsonify({"ok": True, "lesson": lesson.to_dict()})


@api_bp.delete("/lessons/<int:lesson_id>")
def delete_lesson(lesson_id):
    require_admin()
    lesson = Lesson.query.get_or_404(lesson_id)
    old_data = lesson.to_dict()
    Participant.query.filter_by(lesson_id=lesson.id).delete()
    db.session.delete(lesson)
    db.session.commit()
    notify_lesson_change("deleted", lesson, old=old_data)
    return jsonify({"ok": True})

@api_bp.delete("/lessons")
def clear_lessons():
    require_admin()
    db.session.execute(db.text("DELETE FROM participants"))
    db.session.execute(db.text("DELETE FROM lessons"))
    db.session.commit()
    return jsonify({"ok": True})

@api_bp.get("/subjects")
def get_subjects():
    subjs = db.session.query(Lesson.subject).filter(Lesson.subject.isnot(None)).distinct().all()
    return jsonify([s[0] for s in subjs if s[0]])


@api_bp.get("/export_week")
def export_week():
    require_admin()
    start_str = request.args.get("start")
    end_str = request.args.get("end")
    try:
        start_date = datetime.strptime(start_str, "%Y-%m-%d").date()
        end_date = datetime.strptime(end_str, "%Y-%m-%d").date()
    except Exception:
        return jsonify({"ok": False, "error": "Некорректные даты"}), 400

    lessons = Lesson.query.filter(
        Lesson.lesson_date >= start_date,
        Lesson.lesson_date <= end_date
    ).order_by(Lesson.lesson_date, Lesson.pair).all()

    import io
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["Дата", "Пара", "Дисциплина", "Преподаватель", "Аудитория", "Статус", "Комментарий"])
    for l in lessons:
        ws.append([
            l.lesson_date.strftime("%Y-%m-%d"),
            l.pair,
            l.subject,
            l.teacher,
            l.audience or "",
            l.status,
            l.comment or ""
        ])

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)

    from flask import send_file
    return send_file(
        bio,
        as_attachment=True,
        download_name="schedule_week.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

@api_bp.post("/import_excel")
def import_excel():
    require_admin()
    if "file" not in request.files:
        return jsonify({"ok": False, "error": "Файл не передан"}), 400

    file = request.files["file"]
    if not file.filename:
        return jsonify({"ok": False, "error": "Пустое имя файла"}), 400

    try:
        wb = load_workbook(file, data_only=True)
        sheet = wb.worksheets[0]
    except Exception:
        return jsonify({"ok": False, "error": "Не удалось прочитать Excel"}), 400

    day_map = {
        "пн": 0, "по": 0, "понедельник": 0,
        "вт": 1, "во": 1, "вторник": 1,
        "ср": 2, "ср": 2, "среда": 2,
        "чт": 3, "че": 3, "четверг": 3,
        "пт": 4, "пя": 4, "пятница": 4,
        "сб": 5, "су": 5, "суббота": 5,
    }

    def resolve_weekday(raw: str):
        if not raw:
            return None
        norm = re.sub(r"[^а-яa-z]", "", raw.lower())
        for key, val in day_map.items():
            if norm.startswith(key):
                return val
        return None

    def dates_for_weekday(weekday: int):
        year = date.today().year
        start = date(year, 9, 1)
        end = date(year, 12, 30)
        current = start
        while current <= end:
            if current.weekday() == weekday:
                yield current
            current += timedelta(days=1)

    created = 0
    current_day = None
    current_dates = []

    for row in sheet.iter_rows(min_row=10):
        day_cell = row[0].value
        if day_cell:
            wd = resolve_weekday(str(day_cell))
            if wd is not None:
                current_day = wd
                current_dates = list(dates_for_weekday(wd))
        if current_day is None or not current_dates:
            continue

        pair_val = row[1].value
        if pair_val is None or str(pair_val).strip() == "":
            continue
        try:
            pair_num = int(re.sub(r"[^\d]", "", str(pair_val)))
        except Exception:
            continue

        teacher_cell = str(row[6].value).strip() if row[6].value else ""
        if not teacher_cell:
            continue
        teacher_candidates = re.split(r"[;/]", teacher_cell)
        teacher_candidates = [t.strip() for t in teacher_candidates if t.strip()]
        if not teacher_candidates:
            teacher_candidates = [teacher_cell]

        subjects_cells = [c for c in row[2:6] if c.value]
        for cell in subjects_cells:
            subject_chunks = [s.strip() for s in str(cell.value).split("//") if s and s.strip()]
            for idx, subject_raw in enumerate(subject_chunks):
                audience = None
                if "," in subject_raw:
                    subj_split = subject_raw.rsplit(",", 1)
                    if len(subj_split) == 2:
                        maybe_aud = subj_split[1].strip()
                        if maybe_aud:
                            audience = maybe_aud
                        subject_raw = subj_split[0].strip()

                teacher = teacher_candidates[idx] if idx < len(teacher_candidates) else teacher_candidates[0]

                for d in current_dates:
                    lesson = Lesson(
                        lesson_date=d,
                        pair=pair_num,
                        subject=subject_raw,
                        teacher=teacher,
                        audience=audience,
                        status="scheduled",
                    )
                    db.session.add(lesson)
                    created += 1

    db.session.commit()
    return jsonify({"ok": True, "created": created})


@api_bp.post("/notify")
def api_notify():
    data = request.get_json(silent=True) or {}
    user_id = data.get("user_id")
    message = data.get("message")
    if not user_id or not message:
        return jsonify({"ok": False, "error": "user_id и message обязательны"}), 400
    try:
        send_message(int(user_id), str(message))
    except Exception:
        return jsonify({"ok": False, "error": "Не удалось отправить"}), 500
    return jsonify({"ok": True})


@api_bp.post("/subscribers")
def add_subscriber():
    data = request.get_json(silent=True) or {}
    chat_id = str(data.get("chat_id") or "").strip()
    username = (data.get("username") or "").strip()
    full_name = (data.get("full_name") or "").strip()
    if not chat_id:
        return jsonify({"ok": False, "error": "chat_id обязателен"}), 400
    found = Subscriber.query.filter_by(chat_id=chat_id).first()
    if not found:
        found = Subscriber(chat_id=chat_id, username=username, full_name=full_name)
        db.session.add(found)
    else:
        found.username = username or found.username
        found.full_name = full_name or found.full_name
    db.session.commit()
    return jsonify({"ok": True, "id": found.id})


@api_bp.get("/subscribers_list")
def list_subscribers():
    subs = Subscriber.query.all()
    return jsonify([{ "chat_id": s.chat_id, "username": s.username, "full_name": s.full_name } for s in subs])


@api_bp.delete("/subscribers/<chat_id>")
def remove_subscriber(chat_id):
    sub = Subscriber.query.filter_by(chat_id=str(chat_id)).first()
    if sub:
        db.session.delete(sub)
        db.session.commit()
    return jsonify({"ok": True})


@api_bp.get("/lessons/<int:lesson_id>/participants")
def get_participants(lesson_id):
    lesson = Lesson.query.get_or_404(lesson_id)
    parts = Participant.query.filter_by(lesson_id=lesson.id).all()
    return jsonify([{"id": p.id, "name": p.name, "contact": p.contact} for p in parts])


@api_bp.post("/lessons/<int:lesson_id>/participants")
def add_participant(lesson_id):
    lesson = Lesson.query.get_or_404(lesson_id)
    data = request.get_json(silent=True) or {}
    name = (data.get("name") or "").strip()
    contact = (data.get("contact") or "").strip()
    if not name or not contact:
        return jsonify({"ok": False, "error": "Имя и контакт обязательны"}), 400
    p = Participant(lesson_id=lesson.id, name=name, contact=contact)
    db.session.add(p)
    db.session.commit()
    return jsonify({"ok": True, "participant": {"id": p.id, "name": p.name, "contact": p.contact}})

@api_bp.post("/notify")
def manual_notify():
    data = request.json
    telegram_id = data.get("telegram_id")
    text = data.get("text")

    from bot.main import notify_user 

    notify_user(telegram_id, text)

    return jsonify({"ok": True})
